from pathlib import Path
from shutil import copy2
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
ASSETS=ROOT/'comparison-assets'; ASSETS.mkdir(exist_ok=True)
OUT=ROOT/'Experiment 8.docx'
BACKUP=ROOT/'qa'/'Experiment 8 previous analysis.docx'
if OUT.exists() and not BACKUP.exists():
    BACKUP.parent.mkdir(exist_ok=True); copy2(OUT,BACKUP)
items=[
('Overview','06-overview.jpg','overview.jpg',
 'Large empty charts dominated the overview without a clear first action.',
 'A first expense panel and Record expense button guide the user; empty charts take less space.',
 'Recognition and visual hierarchy make the next task obvious. A reusable NextStep component keeps first use guidance consistent.'),
('Transactions','07-transactions.jpg','transactions.jpg',
 'An empty account was described as having no transactions matching the filters.',
 'The empty ledger now explains how to begin and offers Record first expense. Filter guidance is explicit.',
 'Visibility of system status distinguishes an empty ledger from a filtered search with no matches.'),
('Expense entry','08-transaction-modal.jpg','transaction-form.jpg',
 'The form presented payment and classification fields without a clear review sequence.',
 'Numbered sections separate payment details from classification review. A zero amount shows an inline error.',
 'Error prevention gives immediate, field specific feedback. Shared Input and Modal components encapsulate accessible error and keyboard behavior.'),
('Behaviour insights','13-insights-behaviour.jpg','behaviour.jpg',
 'Positive spending messages appeared even though no transactions had been analyzed.',
 'A waiting for your first record state explains the missing evidence and links to expense entry.',
 'Visibility of system status avoids presenting missing analysis as a positive conclusion.'),
('Income insights','14-insights-income.jpg','income.jpg',
 'The safe income label displayed the fractional percentile as 0.25th percentile.',
 'The label shows 25th percentile and a short explanation describes what the figure means.',
 'Abstraction normalizes the API fraction before display. Consistency and plain language keep the label and explanation aligned.'),
('Retirement planning','18-plan-fire.jpg','retirement.jpg',
 'Missing planning inputs produced a Fully Funded contribution status and zero valued targets.',
 'Missing targets show Not set or Not calculated. The contribution panel requests the details needed for a calculation.',
 'Explicit state modeling and error prevention separate unavailable calculations from a valid zero contribution gap.'),
('Scenario inputs','19-plan-scenario.jpg','scenario.jpg',
 'Sliders made precise assumptions difficult to enter and compare.',
 'Each main assumption pairs a slider with an exact numeric field. The monthly investment example is 12750.',
 'User control supports both quick adjustment and precise entry. The reusable RangeField keeps values, bounds, labels, and validation together.'),
('Report exports','23-reports-export.jpg','exports.jpg',
 'A large technical security explanation occupied space above the download choices.',
 'Download choices move higher. Format and spreadsheet safety details remain available in an expandable section.',
 'Progressive disclosure prioritizes the main task while retaining supporting technical information.'),
('Family setup','24-family.jpg','family.jpg',
 'The setup card described household features without a visible invitation sequence.',
 'Three visible steps explain invitation, acceptance, and the shared household view.',
 'Recognition reduces the need to infer the workflow. Clear sequencing explains what happens after an invitation.'),
('Profile','25-profile.jpg','profile.jpg',
 'The profile route mainly displayed planned features and roadmap content.',
 'The page shows current account details and saved financial baseline, with a working Family shortcut. Future features are collapsed.',
 'Visibility of system status exposes useful current information. Progressive disclosure gives planned features less prominence.')
]

def crop(src,dst,modal=False):
    im=Image.open(src).convert('RGB'); w,h=im.size
    # Preserve the same workspace area in each pair; remove the repeated sidebar.
    box=(round(w*.195),round(h*.073),round(w*.98),h)
    if modal:
        # Keep the entire dialog and its dimmed application context visible.
        box=(round(w*.25),round(h*.073),round(w*.80),h)
    im.crop(box).save(dst,quality=96)

D=Document(); s=D.sections[0]
s.page_width=Inches(11.69); s.page_height=Inches(8.27)
s.top_margin=Inches(.43); s.bottom_margin=Inches(.43)
s.left_margin=s.right_margin=Inches(.45)
s.header_distance=s.footer_distance=Inches(.18)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=D.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
for st in D.styles:
    for b in st.element.xpath('.//w:pBdr'): b.getparent().remove(b)
D.styles['Normal'].font.size=Pt(10)
D.styles['Normal'].paragraph_format.space_after=Pt(6)
D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(23); D.styles['Title'].paragraph_format.space_after=Pt(3)
D.styles['Heading 1'].font.size=Pt(16)
D.styles['Heading 1'].paragraph_format.space_before=Pt(0)
D.styles['Heading 1'].paragraph_format.space_after=Pt(8)
D.core_properties.title='Experiment 8'
D.core_properties.subject='FINAURA frontend before and after software design improvements'
D.core_properties.author='FINAURA Project Team'
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('FINAURA  |  Experiment 8  |  '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def text(label,body):
    p=D.add_paragraph(); p.add_run(label+' ').bold=True; p.add_run(body)
    return p

def border(cell,header=False):
    pr=cell._tc.get_or_add_tcPr(); bs=OxmlElement('w:tcBorders')
    for e in ['top','left','bottom','right']:
        b=OxmlElement('w:'+e); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); bs.append(b)
    pr.append(bs)
    margins=OxmlElement('w:tcMar')
    for edge,value in [('top','90'),('bottom','90'),('left','85'),('right','85')]:
        x=OxmlElement('w:'+edge); x.set(qn('w:w'),value); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    if header:
        shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E8EEF8'); pr.append(shade)

D.add_paragraph('Experiment 8','Title')
p=D.add_paragraph('FINAURA frontend improvements through software design principles')
p.paragraph_format.space_after=Pt(10)
D.add_heading('Hypothetical user feedback',1)
D.add_paragraph('These fabricated comments illustrate possible feedback that could motivate the UI changes. They are not collected participant responses. Person numbers match the numbered comparisons that follow.')
feedback=[
    "The overview is mostly empty charts. I cannot tell where to start or how to record my first expense.",
    "I have not added any transactions, so why does the page say none match my filters? Please explain how to begin.",
    "I am unsure which expense details to enter first and when to check the category. Please tell me beside the amount field if zero is invalid.",
    "I have not recorded any spending yet. How can the app say my spending is disciplined? Show when there is too little data to analyze.",
    "What does 0.25th percentile mean? I need a clear label and a short explanation of how the reliable income figure should be read.",
    "I have not completed my retirement details, but the contribution panel says Fully Funded. Please show when a calculation is not available.",
    "I want to try an exact monthly investment of 12750. The sliders make precise values difficult to enter and compare.",
    "I just want to download my records. The large technical explanation distracts me from the download choices; make those details optional.",
    "What happens after I invite someone to my family? Please show whether they must accept and when the shared household view becomes available.",
    "The Profile page mainly shows future features. I want to see my current account details and saved financial information first."
]
t=D.add_table(rows=1,cols=2); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
widths=[1.25,9.54]
for col,width in zip(t.columns,widths): col.width=Inches(width)
for cell,label in zip(t.rows[0].cells,['Person','Hypothetical feedback']): cell.text=label
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
for i,comment in enumerate(feedback,1):
    cells=t.add_row().cells; cells[0].text=f'Person {i}'; cells[1].text=comment
for ri,row in enumerate(t.rows):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for ci,cell in enumerate(row.cells):
        cell.width=Inches(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        border(cell)
        if ri==0:
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F2F2F2'); cell._tc.get_or_add_tcPr().append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.1
            if ci==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:r.font.size=Pt(11);r.bold=ri==0

for idx,(title,old,new,before,after,principle) in enumerate(items):
    D.add_page_break()
    if idx==0:
        p=D.add_paragraph('Original UI on the left and implemented live UI on the right. Screenshots use the same local test account and are cropped to the relevant workspace.')
        p.paragraph_format.space_after=Pt(9)
        for r in p.runs:r.font.size=Pt(9)
    D.add_heading(f'{idx+1} {title}',1)
    oldp=ASSETS/f'{idx+1:02d}-before.jpg'; newp=ASSETS/f'{idx+1:02d}-after.jpg'
    crop(ROOT/'screenshots'/old,oldp,idx==2); crop(ROOT/'after'/new,newp,idx==2)
    t=D.add_table(rows=2,cols=2); t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c in t.columns:c.width=Inches(5.395)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(5.395); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border(cell,ri==0)
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1
            if ri==0:
                r=p.add_run('BEFORE   ORIGINAL UI' if ci==0 else 'AFTER   UPDATED LIVE UI'); r.bold=True; r.font.size=Pt(10)
            else:
                f=oldp if ci==0 else newp
                im=Image.open(f); width=5.20; height=width*im.height/im.width
                if height>4.10: height=4.10; width=height*im.width/im.height
                p.add_run().add_picture(str(f),width=Inches(width),height=Inches(height))
                p._p.xpath('.//wp:docPr')[0].set('descr',f'{title} '+('before changes' if ci==0 else 'after implemented changes'))
    p=D.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=.3;p.add_run().font.size=Pt(2)
    text('Before',before)
    text('After',after)
    text('Design principle',principle)
D.save(OUT)
print(OUT)
