from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
SHOTS=ROOT/'screenshots'
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
sec.top_margin=Inches(.58); sec.bottom_margin=Inches(.58)
sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
sec.footer_distance=Inches(.25)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[n]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
# Remove decorative borders inherited from the bundled Word default styles.
for s in doc.styles:
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
for n,sz in [('Title',30),('Subtitle',14),('Heading 1',17),('Heading 2',12),('Heading 3',11)]:
    doc.styles[n].font.size=Pt(sz)
    doc.styles[n].paragraph_format.space_before=Pt(11)
    doc.styles[n].paragraph_format.space_after=Pt(7)
doc.styles['Caption'].font.size=Pt(8.5)
doc.styles['Caption'].font.italic=False
doc.styles['Caption'].paragraph_format.space_after=Pt(7)
doc.styles['Caption'].paragraph_format.line_spacing=1.05
doc.core_properties.title='Experiment 8'
doc.core_properties.subject='FINAURA frontend analysis and software design principles'
doc.core_properties.author='FINAURA Project Team'

def p(text,style=None,bold=False):
    par=doc.add_paragraph(style=style); r=par.add_run(text); r.bold=bold
    return par
def h(text,level=1): return doc.add_heading(text,level)
def page(title): doc.add_page_break(); h(title)
def cap(text): return p(text,'Caption')
def shot(file,label,width=6.35):
    par=doc.add_paragraph(); par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after=Pt(3)
    par.add_run().add_picture(str(SHOTS/file),width=Inches(width))
    # Alt text describes the evidence rather than the file.
    par._p.xpath('.//wp:docPr')[0].set('descr',label)
    cap(label)
def grid(items,width=3.17):
    # Figure comparisons use a consistent grid with light cell borders.
    t=doc.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col in t.columns: col.width=Inches(3.43)
    for i in range(0,len(items),2):
        cells=t.add_row().cells
        for j,item in enumerate(items[i:i+2]):
            f,label=item; c=cells[j]; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            par=c.paragraphs[0]; par.paragraph_format.space_after=Pt(3)
            par.add_run().add_picture(str(SHOTS/f),width=Inches(width))
            par._p.xpath('.//wp:docPr')[0].set('descr',label)
            para=c.add_paragraph(label,style='Caption')
        trPr=t.rows[-1]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
        for c in cells:
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            c._tc.get_or_add_tcPr().append(borders)
    p('').paragraph_format.space_after=Pt(0)
def table(headers,rows,widths,size=9):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,txt in zip(cells,row): c.text=str(txt)
    for ri,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for e in ['top','left','bottom','right']:
                x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for e,v in [('top','85'),('bottom','85'),('left','95'),('right','95')]:
                x=OxmlElement('w:'+e); x.set(qn('w:w'),v); x.set(qn('w:type'),'dxa'); margins.append(x)
            tcPr.append(margins)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'17365D' if ri==0 else ('F2F5F9' if ri%2==0 else 'FFFFFF')); tcPr.append(shade)
            for par in c.paragraphs:
                par.paragraph_format.space_after=Pt(0); par.paragraph_format.line_spacing=1.08
                if ci==0 and len(headers)>3: par.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size=Pt(size)
                    if ri==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    p('').paragraph_format.space_after=Pt(0)
    return t

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('FINAURA  |  Experiment 8  |  '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

p('Experiment 8','Title')
p('Frontend analysis using software design principles','Subtitle')
p('FINAURA Adaptive Financial Intelligence System',bold=True)
p('Web interface evaluation and mobile frontend source review\nEvaluation period 15 and 16 October 2026  |  Repository revision 76d674b8')
h('Aim')
p('To evaluate the existing FINAURA frontend, identify design problems with observable evidence, and propose changes that can be traced to software design and human interface principles. The comparison covers both what users see and how frontend responsibilities are organized.')
h('Main result')
p('FINAURA already has a consistent workspace shell, reusable web components, domain hooks, responsive layouts, and explicit privacy messaging. The highest-priority improvements are to distinguish missing data from positive financial conclusions, correct display contracts, complete keyboard interaction, and make validation errors actionable. Targeted changes can preserve the current visual identity while improving clarity and maintainability.')
h('Application context')
p('The React web portal organizes personal finance tasks into Overview, Activity, Insights, Plan, Reports, and Family. The Expo mobile client supports daily entry, financial diagnostics, planning, household views, and an assistant. Both clients consume the Express API and the financial modeling services.')
h('Method')
p('The local web app, API, and ML service were started on the computer. Public routes and every principal authenticated workspace tab were inspected in the browser. Screenshots were captured from the running application. Source inspection covered routes, page composition, shared controls, hooks, formatting, mobile screens, navigation, and stores. Keyboard and reload checks supplemented the visual review.')
h('How to read the comparison')
p('Figures marked Existing are live app captures. Figures marked Proposed are separate design mockups and do not represent deployed changes. Proposed input examples are explicitly labeled; financial outcomes are not presented as measured improvements. D01 to D18 identify recommendations throughout the document and connect them to the traceability matrices.')

page('Design principles and review coverage')
p('Software architecture principles guide where a change belongs. Interface principles guide how the resulting behavior should work for users. The interface review draws on Nielsen’s heuristics [1] and W3C interaction and accessibility guidance [2–4].')
table(['Principle','Application to FINAURA'],[
('Single responsibility and cohesion','Keep screen rendering, input validation, API orchestration, and financial interpretation in focused units.'),
('Abstraction and encapsulation','Expose stable UI contracts for fields, dialogs, tabs, request states, and view models.'),
('DRY and consistency','Centralize vocabulary, formatters, status semantics, and visual tokens instead of repeating policy.'),
('Low coupling and dependency inversion','Let screens depend on typed domain interfaces; keep networking and response normalization behind adapters.'),
('Open closed design','Allow new workspace tabs and status variants through configuration and composition without editing every consumer.'),
('Feedback and error prevention','Show current request state, specific errors, and data sufficiency before making a conclusion.'),
('Recognition and progressive disclosure','Make common tasks obvious and place advanced assumptions behind clear optional controls.'),
('Accessibility and user control','Support keyboard use, readable text, linked field errors, explicit review, and predictable focus.')
],[2.1,4.77],9.5)
h('Coverage inventory',2)
table(['Area','Reviewed surfaces'],[
('Public web','Landing, login, registration, forgot password, reset password, and 404 recovery.'),
('Activity','Transactions, income, liabilities, and the transaction entry modal.'),
('Insights','Spending, FMI health, behaviour, and income analysis.'),
('Planning','Net worth, assets, goals, FIRE, and Scenario Lab.'),
('Other web','Overview, all four report tabs, Family, Profile, Settings, and narrow-screen navigation.'),
('Mobile source','All 22 screen files, the tab and stack navigators, authentication flow, shared components, API layer, and stores.')
],[1.2,5.67],9.3)
p('Evidence conditions: the supplied local test account had a ₹20,000 balance and no transaction, income, asset, goal, liability, or family records. This provides direct evidence of first-use states. Populated chart interactions and active household membership states were assessed from source rather than represented as live tested states.',style='Caption')

page('Public pages and account access')
grid([
('02-login.jpg','Figure 1  Existing login with labeled fields and a password visibility control.'),
('03-register.jpg','Figure 2  Existing registration with a visible password checklist.'),
('04-forgot-password.jpg','Figure 3  Existing recovery entry page.'),
('05-reset-password.jpg','Figure 4  Existing reset page asks for a token and new password.')
],2.95)
p('D03  Keep the current labels and password visibility controls. Add field-specific errors connected with aria-describedby and aria-invalid. Use the same password policy in registration and reset. Login should surface the server field error instead of the generic “Validation failed” message. A reset link should prefill its token when supported by the recovery contract; retain a manual token fallback.')
cap('Source evidence  web/src/components/ui/Input.tsx; web/src/pages/LoginPage.tsx; RegisterPage.tsx; ResetPasswordPage.tsx; server/validators/authValidator.js. No registration or password reset was submitted during this review.')

page('Landing page and first use')
shot('01-landing.jpg','Figure 4A  Existing live landing page. The dashboard visible within this page is explicitly labeled Illustrative Preview by the app itself.',4.25)
p('The landing page has clear Sign In and Get Started actions and separates the mobile and web experiences farther down the page. Its preview label usefully prevents sample financial figures from appearing to be account data. Preserve those strengths.')
table(['Existing language','Proposed difference','Principle'],[
('Single-Count Accounting and Deterministic Modeling are prominent trust labels.','Explain the user benefit in plain language, such as avoiding duplicate records and seeing the assumptions behind projections.','D15 Match users’ language and progressively disclose detail.'),
('Long technical sections explain the product architecture.','Lead with three concrete user tasks and short examples; retain deeper explanations below.','D08 Recognition and task hierarchy.'),
('FMI expands to Financial Momentum Index.','Use the canonical glossary consistently in the landing page, signed-in app, and reports.','D01 Consistency and DRY.')
],[2.1,2.95,1.82],9.4)
cap('Source  web/src/pages/LandingPage.tsx. The landing preview is part of the existing application and is distinct from the Proposed mockups in this report.')

page('Overview comparison')
shot('06-overview.jpg','Figure 5  Existing overview. Four summary cards and large empty charts dominate the first viewport.',6.05)
shot('proposed-overview.jpg','Figure 6  Proposed overview. D01 and D08 add explicit period context and a first expense action; D04 strengthens supporting text. Mockup uses the same balance and FMI example as the live account.',6.05)

page('Transaction workspace comparison')
shot('07-transactions.jpg','Figure 7  Existing transactions. No entries exist, but the message describes a failed filter match.',6.05)
shot('proposed-activity.jpg','Figure 8  Proposed transactions. D05 retains filter state; D08 separates an empty ledger from zero search results. The active period, matching count, and first entry action remain visible. Proposed layout, not an implemented result.',6.05)

page('Transaction entry comparison')
shot('08-transaction-modal.jpg','Figure 9  Existing transaction dialog. It has labels, Escape handling, and an autofocus amount field, but lacks a complete focus cycle.',6.05)
shot('proposed-form.jpg','Figure 10  Proposed entry form. D02 defines the focus contract; D03 places errors beside fields; D06 makes classification a deliberate review action. This mockup illustrates behavior that must be implemented and tested.',6.05)

page('Income and recurring liabilities')
shot('09-income.jpg','Figure 11  Existing income tab. It already explains its empty state and offers Record Income.',5.95)
shot('10-liabilities.jpg','Figure 12  Existing liabilities tab. It already offers a first liability action. D01 proposes the navigation label Recurring bills with the technical term explained in context; retain due dates, payment status, and auto-deduction details.',5.95)
cap('D08  Preserve these direct empty-state actions and apply the same pattern to Overview and Transactions. D02 and D03 should be inherited by income and liability dialogs through shared primitives.')

page('Insights and truthful interpretation')
grid([
('11-insights-spending.jpg','Figure 13  Existing spending analysis.'),
('12-insights-fmi.jpg','Figure 14  Existing FMI health view.'),
('13-insights-behaviour.jpg','Figure 15  Existing behaviour view with zero analyzed records.'),
('14-insights-income.jpg','Figure 16  Existing income metrics with a 0.25th percentile label.')
])
h('D09 Separate evidence from a positive conclusion',2)
p('The Behaviour tab reports zero analyzed records and then says recent spending shows disciplined consistency. Replace that state with “Not enough transaction history to assess spending behaviour” and an action to record expenses. Show a positive no-pattern message only when a defined evidence threshold is satisfied. FMI can still display its backend score, but should include the observation count and explain limited evidence.')
h('D16 Normalize units at the presentation boundary',2)
p('The API represents the reliable-income percentile as a fraction. The frontend inserts 0.25 directly into an ordinal label. Convert a documented percentile fraction to a display value of 25 and render “25th percentile.” Keep this conversion in a typed adapter so all income views use the same unit.')
cap('Evidence  BehaviorSignalCards.tsx:52; useBehaviorInsights.ts:31; useIncomeInsights.ts:138; server/utils/incomeAnalytics.js:204 and 615. These are observed interpretation defects, not measured model accuracy failures.')

page('Net worth assets and goals')
grid([
('15-plan-net-worth.jpg','Figure 17  Existing net worth and balance-sheet overview.'),
('16-plan-assets.jpg','Figure 18  Existing asset ledger with a first asset action.')
])
shot('17-plan-goals.jpg','Figure 19  Existing goals view. The summary and Create Goal action establish a useful reusable pattern.',5.95)
p('D08  For an empty balance sheet, replace the all-zero composition chart with a short explanation and links to add assets or liabilities. D14  Load the active planning domain when needed and share common snapshots; the current page mounts all five domain hooks even when only one tab is visible.')
cap('Source  web/src/pages/PlanPage.tsx:69; hooks/useNetWorth.ts; useAssets.ts; useGoals.ts; useFirePlan.ts; useScenarioLab.ts. Preserve the existing distinction between operating balance and recorded asset-based net worth.')

page('Scenario Lab comparison')
shot('19-plan-scenario.jpg','Figure 20  Existing Scenario Lab. Monthly investment uses a slider with a ₹1,000 step, and current outputs are unavailable.',6.05)
shot('proposed-scenario.jpg','Figure 21  Proposed Scenario Lab. D07 adds exact numeric entry beside sliders; D09 and D14 explicitly show missing profile data and an uncalculated result. ₹12,750 is an illustrative input, not a simulated outcome.',6.05)

page('Retirement and data sufficiency')
shot('18-plan-fire.jpg','Figure 22  Existing FIRE view. An unavailable required investment appears beside a ₹0 Fully Funded contribution-gap message.',6.35)
h('D09 Use an explicit unavailable state',2)
p('FireContributionPanel decides whether a gap exists with “gap is non-null and positive.” The alternative includes both a real zero and an unknown value, yet displays “₹0 (Fully Funded).” The live account demonstrates the conflict: required investment is unavailable while the gap card reports full funding.')
table(['Existing behavior','Proposed behavior','Principle'],[
('Null or non-positive gap follows the success branch.','Use unavailable, funded, and shortfall as separate states.','Explicit state modeling and error prevention.'),
('A zero target can produce a positive buffer statement.','Require valid essential-expense and target inputs before asserting coverage.','Feedback based on sufficient evidence.'),
('Technical fallback explanation dominates the lower view.','State which calculation is unavailable, what remains usable, and the next step.','Visibility of system status and plain language.')
],[2.15,2.95,1.77],9.4)
p('Acceptance: null requirement or an unresolved target never produces Fully Funded. A genuine zero gap with a valid target can still show that status. These changes affect presentation state; the deterministic formulas remain in the financial domain service.')
cap('Source  web/src/components/plan/FireContributionPanel.tsx:20 and 53; FireHeroSummary.tsx; EmergencyFundPanel.tsx; MonteCarloPanel.tsx.')

page('Reports and export clarity')
grid([
('20-reports-weekly.jpg','Figure 23  Existing rolling weekly review.'),
('21-reports-monthly.jpg','Figure 24  Existing calendar month review.'),
('22-reports-history.jpg','Figure 25  Existing history empty state.'),
('23-reports-export.jpg','Figure 26  Existing export workspace.')
])
p('The report tabs already distinguish weekly and calendar-month scopes, provide print actions, and expose record counts. History explains how data becomes available. Keep these strengths and carry period context into every chart, file name, and export preview.')
table(['Change','Existing','Proposed'],[
('D01 Vocabulary','The monthly report uses Financial Mind Index while Insights uses Financial Momentum Index.','Choose one FMI expansion, preferably the documented Financial Maturity Index, and centralize it.'),
('D09 Evidence','A no-activity week still shows optimal health and healthy spending language.','Explain no activity or limited evidence before interpreting behavior.'),
('D15 Progressive disclosure','RFC 4180, UTF-8 BOM, and formula characters precede export tasks.','Lead with data type, record count, period, and download. Put technical format details under an optional help control.')
],[.9,2.6,3.37],9.3)
cap('Source  components/reports/MonthlyReportView.tsx:458; WeeklyReportView.tsx; ExportReportView.tsx. CSV formula sanitization should remain active; only its presentation changes.')

page('Family and household privacy')
shot('24-family.jpg','Figure 27  Existing no-family state. Privacy and the invitation action are both clearly explained.',6.4)
p('Preserve the separation of personal accounts and household aggregates. FamilyPrivacyBanner and NoFamilyCard already express this boundary. The accepted-invitation creation rule and six-member limit help explain the next step. Do not add personal transaction drill-downs as a visual enhancement.')
table(['Visible in household views','Remain personal'],[
('Pooled income and spending totals','Individual transaction descriptions and income sources'),
('Household FMI and weighted pillars','Personal balances, assets, liabilities, and personal FMI'),
('Member roles and invitation status','Individual retirement targets and assumptions')
],[3.43,3.44],10)
p('D02 and D03  Apply the shared accessible dialog and validation contract to invitation forms. D14  Represent invitation and membership requests with clear loading, failure, and retry states. Keep owner-only controls conditional on the authenticated role.')
cap('Live coverage is the no-family state. Active household, invitation, role, removal, and pooled-display behavior was reviewed in FamilyPage.tsx, useFamily.ts, family components, and the existing family web tests. No invitations were sent or membership changes made.')

page('Availability and recovery')
grid([
('25-profile.jpg','Figure 28  Existing Profile is a roadmap placeholder.'),
('26-settings.jpg','Figure 29  Existing Settings is a roadmap placeholder.')
])
shot('27-not-found.jpg','Figure 30  Existing 404 gives routes back to Overview and the public home page.',4.8)
table(['Change','Existing','Proposed'],[
('D10 Availability','Profile and Settings are full navigation destinations but contain planned capabilities.','Mark availability in navigation and show usable account information or a concise unavailable state with a working next step.'),
('D15 User language','Settings exposes sessionStorage policy and backend infrastructure terms.','Use Session status, Account security, and Service availability; place implementation details in developer documentation.'),
('D11 Semantics','404 links wrap Button elements, creating nested interactive markup.','Render each recovery action as one styled link with one focus target.')
],[.9,2.5,3.47],9.3)
cap('Source  web/src/App.tsx:134 and 161; pages/PlaceholderPage.tsx; pages/NotFoundPage.tsx. Profile editing already exists in the mobile source and provides a useful model for web feature parity.')

page('Responsive layout and keyboard control')
grid([
('28-mobile-web-overview.png','Figure 31  Existing web overview at 390 CSS pixels. Summary cards stack into a single column.'),
('29-mobile-web-drawer.png','Figure 32  Existing mobile web drawer. This is the web portal, not the Expo mobile client.')
],2.4)
p('The narrow overview fits the viewport, and the drawer provides access to the same major destinations. Preserve this layout. D11 should add aria-expanded and aria-controls to the trigger, move focus into an opened drawer, contain interaction while it is modal, and restore focus when it closes.')
p('D02 was verified in the transaction form: pressing Tab from its final submit button moved focus outside the dialog to the body; closing the dialog also left focus on the body. D11 was verified in Insights: ArrowRight on the first tab did not advance focus or selection. Use one shared keyboard tab pattern with roving tabIndex, linked tab panels, arrow keys, Home, and End.')
cap('Sources  layouts/AppLayout.tsx; components/layout/Sidebar.tsx; TopBar.tsx; components/ui/Modal.tsx; the Activity, Insights, Plan, and Report tab components. W3C APG dialog and tabs patterns define the expected behavior [2, 3].')

changes=[
('D01','FMI has Momentum, Mind, and Maturity expansions.','Consistency and DRY','One glossary for names, units, and period labels.','Medium'),
('D02','Tab can leave the dialog; closing loses trigger focus.','Encapsulation and accessibility','Shared dialog manages initial focus, trapping, Escape, and restoration.','High'),
('D03','Generic login errors; reset checks only password length and match locally.','Error prevention and reusable policy','Shared validation schema, field error IDs, and specific recovery messages.','High'),
('D04','Tertiary text token #94a3b8 has 2.56:1 contrast on white.','Accessibility and consistent tokens','Use a darker semantic text token for normal-sized explanatory text.','High'),
('D05','Transaction search groceries clears on reload.','User control and state ownership','Serialize filters, sort, and page in the URL, preserving tab parameters.','Medium'),
('D06','Source risk: classifier responses auto-apply without an in-flight response guard.','Error prevention and separation of concerns','Extract a classification hook, reject stale responses, and offer explicit review.','High'),
('D07','Monthly scenario input advances in ₹1,000 slider steps.','Efficiency and user control','Pair ranges with numeric fields, bounds, units, and one source of input state.','Medium'),
('D08','Empty transactions use filter-failure copy; overview emphasizes empty charts.','Recognition and task hierarchy','Separate empty ledger from no matches; give one obvious first entry action.','Medium'),
('D09','Zero evidence can yield disciplined, optimal, or Fully Funded wording.','Explicit state modeling and feedback','Model insufficient data separately from measured zero and successful funding.','High'),
('D10','Profile and Settings lead to roadmap placeholders.','Predictability and visibility','Expose availability before navigation; provide useful current account information.','Medium'),
('D11','Tabs lack arrow navigation; drawer trigger lacks expanded state.','Reusable abstraction and accessibility','One configurable tab primitive and one complete drawer focus contract.','High'),
('D12','Large mobile screens and a monolithic web stylesheet mix responsibilities.','Single responsibility and high cohesion','Extract feature components, hooks, and scoped styles; share semantic tokens.','Medium'),
('D13','Mobile user data is copied between auth and general stores.','Single source of truth and low coupling','Keep one user authority and expose derived selectors to consumers.','Medium'),
('D14','Planning mounts all domains; some mobile errors are only logged.','Dependency inversion and explicit states','Typed adapters, shared request cache, active-tab loading, and visible retry states.','Medium'),
('D15','Export and Settings lead with technical infrastructure details.','Progressive disclosure and user language','Lead with tasks and concise outcomes; retain technical help behind details.','Medium'),
('D16','Percentile fraction 0.25 renders as 0.25th percentile.','Typed contracts and abstraction','Normalize fraction to percent in one view-model adapter.','High'),
('D17','Mobile bottom tab labels disable font scaling.','Accessibility and responsive composition','Allow supported text scaling and test label fit with platform accessibility settings.','High'),
('D18','Four mobile screen files have no imports from current navigation flows.','Information architecture and cohesion','Decide whether to expose those features or remove dormant UI after a feature decision.','Medium')
]
for i,title in [(0,'Traceability matrix for core interactions'),(9,'Traceability matrix for frontend structure')]:
    page(title)
    p('Each recommendation links the observed condition to a principle and a specific change. High means potential misleading interpretation or blocked access; Medium means workflow, consistency, or maintenance cost. Priorities are review judgments, not measured incident rates.')
    table(['ID','Existing evidence','Principle','Proposed difference','Priority'],changes[i:i+9],[.43,2.0,1.35,2.48,.61],9.2)
    if i==0:
        p('D04 measurement: the sRGB relative-luminance contrast calculation gives 2.56:1 for #94a3b8 on #ffffff. #475569 on white gives 7.58:1. WCAG 2.2 requires at least 4.5:1 for ordinary text and 3:1 for large text, subject to its stated exceptions [4]. This is a token-pair finding, not a claim that the entire app has been certified or audited for conformance.')
    else:
        p('D12 size evidence: web/src/index.css has 11,671 lines; DashboardScreen.tsx has 1,936; FinancialOutlookScreen.tsx has 2,592; FamilyScreen.tsx has 1,924. Size alone is not a defect. Here, each reviewed mobile file also combines data requests, presentation decisions, interactions, and styles, providing a concrete reason to split responsibilities.')

page('Applying software architecture principles')
h('D12 Single responsibility and cohesion',2)
p('Retain the web page, component, hook, and service separation already present. Split large mobile screens into domain sections such as RetirementSummary, ContributionStrategy, and SimulationResults, supported by a useFinancialOutlook hook. Split index.css by shared tokens, primitives, and feature styles. Each extraction should have one reason to change rather than merely making files shorter.')
h('D02 D03 and D11 Abstraction and encapsulation',2)
p('A Dialog primitive should own the complete focus lifecycle. A Field primitive should own label, input, helper, and error relationships. A Tabs primitive should own keyboard interaction and ARIA relationships, taking tab descriptors and panel content as inputs. Consumers should supply domain labels and actions rather than reimplement these rules.')
h('D01 and D04 DRY with semantic tokens',2)
p('Use a shared glossary for FMI, FIRE, Need, Want, and Investment. Introduce semantic tokens for primary text, supporting text, error text, surfaces, and focus treatment. The mobile source currently declares #2563EB in IncomeFlowScreen and #3B3BDE in FinancialOutlookScreen. Share semantic roles across platforms while allowing platform-specific rendering.')
h('D13 Single source of truth',2)
p('client/App.tsx copies the authenticated user into useStore. ProfileScreen writes updated user data into both stores. Prefer one canonical user store with derived domain selectors so a profile update cannot leave two client copies out of sync. Keep domain collections in their own stores when their lifecycle differs.')
h('D14 Dependency inversion and low coupling',2)
p('Define typed domain adapters such as FinancialOverviewReader and ScenarioEvaluator. Hooks consume those interfaces; Axios normalization stays in the adapter. Screens render explicit loading, ready, empty, insufficient, and error states. Use a shared request cache or active-tab boundary to avoid independent requests for the same planning snapshot.')
h('Open closed composition',2)
p('The shared tab component can accept a new descriptor without changing focus logic. A domain status renderer can accept a supported variant without duplicating message branches across screens. Apply these extensions where repetition already exists; do not introduce an abstract layer for every one-off component.')
cap('Architecture references  React Thinking in React [5]; web/src/services/api.ts; hooks; pages/PlanPage.tsx; client/App.tsx:14–44; client/src/screens/ProfileScreen.tsx:173–177 and 410–418.')

page('Mobile frontend coverage and recommendations')
p('The mobile review covers every screen file, including screens that are conditionally rendered and screens that are not currently reachable. Findings below are based on source inspection; Figures 31 and 32 depict the responsive web portal.')
table(['Flow','Screen files reviewed','Design application'],[
('Account access','LoginScreen, RegisterScreen, ForgotPasswordScreen, ResetPasswordScreen','D03 shared validation and field feedback; preserve password visibility controls.'),
('Entry lifecycle','OnboardingScreen, WelcomeScreen','Explain required financial inputs and keep onboarding completion separate from authentication.'),
('Home and diagnostics','DashboardScreen, FmiScreen, AlertsScreen','D09 evidence-aware diagnostics; D12 section components; explicit loading, empty, and failure states.'),
('Money entry','TransactionsScreen, TransactionEntryScreen, IncomeFlowScreen, UpdateBalanceScreen','D06 review classification; D03 inline errors; preserve a clear distinction between recording an expense and adjusting a balance.'),
('Planning and debts','LiabilitiesScreen, AssetsScreen, FinancialOutlookScreen, GoalPlannerScreen, EnvelopesScreen','D07 precise inputs; D09 unavailable states; D12 focused domain hooks and reusable cards.'),
('Review and assistant','ReportsScreen, ChatScreen','D14 distinguish history-load failure from no history; preserve user input and offer retry for a failed send.'),
('Account and household','ProfileScreen, FamilyScreen','D13 one user authority; retain household privacy boundaries and role-conditional actions.')
],[1.1,3.0,2.77],9.5)
h('D17 Preserve accessibility settings',2)
p('AppNavigator.tsx:82 sets tabBarAllowFontScaling to false, while labels use a 10-point size. Enable an appropriate scaling strategy, allow enough tab height, and test large text on both platforms. Existing accessibility labels on Home, dashboard actions, profile, assets, and family controls are useful foundations.')
h('D18 Resolve dormant screens deliberately',2)
p('AlertsScreen, ReportsScreen, GoalPlannerScreen, and EnvelopesScreen are present but have no import sites in the current app flow. In contrast, TransactionEntryScreen and UpdateBalanceScreen are embedded by other screens, and the recovery screens are conditionally rendered from Login. Decide which dormant features belong in navigation before adding routes or deleting files.')
cap('Inventory  22 screen files in client/src/screens. Navigation evidence  client/src/navigation/AppNavigator.tsx and client/App.tsx; import references were checked across client/src.')

page('Implementation order and acceptance checks')
p('Implement the high-priority correctness and access changes first. Follow with durable state, common primitives, and then structural refactoring. Acceptance checks below describe the target behavior; they are not claims that proposed changes have already passed.')
table(['Change IDs','Acceptance check'],[
('D09 D16','With zero records, show insufficient evidence. With null contribution requirement, never show Fully Funded. A percentile fraction of 0.25 displays as 25th percentile.'),
('D02 D11','Tab and Shift+Tab stay within an open modal. Escape closes it. Focus returns to the opener. Tab lists support arrows, Home, End, and linked panels.'),
('D03','Registration and reset enforce the same password rules. Each error is associated with its field. Invalid email receives a specific message.'),
('D04 D17','Normal supporting text passes 4.5:1 on its actual background. Mobile tab labels remain usable with enlarged system text.'),
('D05','Set search, category, period, sort, and page; reload and use Back. State remains consistent with the URL and the matching summary.'),
('D06','Delay request A, issue request B, and resolve A last. A cannot overwrite B. Manual category and spend-type choices remain user-controlled.'),
('D07','Enter ₹12,750 exactly; the slider and numeric field agree. Reject invalid age and percentage values with field-level feedback.'),
('D08 D10 D15','A new user can identify the next available action. Placeholder status is visible before navigation. Export count and period are clear.'),
('D12 D13 D14','An edit updates every user consumer through one authority. One tab does not trigger unrelated domain fetches. Errors expose a retry action.'),
('D18','Every intended mobile feature has a reachable route or documented embedded entry point. Dormant files are resolved by a product decision.')
],[1.0,5.87],9.4)
h('Checks completed on the existing app',2)
p('Web TypeScript checking passed. The existing web regression command passed 40 family checks and 8 CSV checks. Live inspection confirmed the filter reset, dialog focus escape and lost restoration, unchanged ArrowRight tab behavior, zero-record positive wording, and the percentile label defect. These checks characterize the current frontend; they do not constitute a usability study or a complete runtime test suite.')
cap('Commands  npm --prefix web run typecheck; npm run test:web. Source checks support structural findings. No app source files were changed for this experiment.')

page('Result conclusion and references')
h('Result',2)
p('The review produced 18 traceable recommendations, live screenshots across all principal web routes and tabs, four proposed visual layouts, and a complete mobile screen inventory. The before-and-after comparison makes each recommendation visible while the matrices explain the principle and expected behavior.')
h('Conclusion',2)
p('Applying software design principles to FINAURA means improving shared behavior as well as presentation. Typed display adapters prevent unit mistakes, explicit states prevent unsupported conclusions, and accessible primitives make correct interaction reusable. Focused hooks and a single user authority reduce coupling. These changes can improve the interface without replacing the working workspace structure, financial calculations, or household privacy model.')
h('Primary design references',2)
refs=[
('[1] Jakob Nielsen','10 Usability Heuristics for User Interface Design','https://www.nngroup.com/articles/ten-usability-heuristics/'),
('[2] W3C WAI','ARIA Authoring Practices Guide Dialog Modal Pattern','https://www.w3.org/WAI/ARIA/apg/patterns/dialog-modal/'),
('[3] W3C WAI','ARIA Authoring Practices Guide Tabs Pattern','https://www.w3.org/WAI/ARIA/apg/patterns/tabs/'),
('[4] W3C','Web Content Accessibility Guidelines 2 2','https://www.w3.org/TR/WCAG22/'),
('[5] React','Thinking in React','https://react.dev/learn/thinking-in-react')]
for author,title,url in refs:
    p(author+'  '+title)
    par=p(url); par.paragraph_format.space_after=Pt(9)
    for run in par.runs: run.font.size=Pt(9)
p('References consulted on 15 October 2026. Principle names are used as design criteria; the recommendations are based on the local application evidence described in this report.')
h('Repository evidence guide',2)
p('Web entry points are in web/src/App.tsx and pages. Shared interaction primitives are in web/src/components/ui; workspace components are grouped by feature. Data orchestration is in web/src/hooks and response handling in web/src/services. Mobile flow starts in client/App.tsx and client/src/navigation/AppNavigator.tsx, with screens and stores under client/src. Source paths and line locators throughout the report refer to revision 76d674b8.')
h('Reproducibility',2)
p('Existing UI captures came from http://127.0.0.1:5173 using the supplied local test account. Desktop captures used a browser viewport of approximately 1365 by 900 pixels; initial public-page captures used the smaller in-app viewport. The narrow-screen evidence uses 390 by 844 CSS pixels. The proposed pages were rendered separately on localhost port 5174. Captures span late 15 October and early 16 October 2026 in India Standard Time.')

output=ROOT/'Experiment 8.docx'
doc.save(output)
print(output)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables),'images',len(doc.inline_shapes))
